"""Export author-review DOCX copies from canonical Markdown for Word coauthor editing."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "reports/robust_v2/submission_review_20261014"


def inline(p, content):
    tokens = re.split(r"(\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*))", content)
    for part in tokens:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*"):
            p.add_run(part[1:-1]).italic = True
        else:
            p.add_run(part)


def table(doc, lines):
    rows = [[c.strip() for c in line.strip().strip("|").split("|")] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r"[-:\s]+", c or "-") for c in r)]
    if not rows:
        return
    n = max(map(len, rows))
    t = doc.add_table(rows=len(rows), cols=n)
    t.style = "Light Shading Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, content in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            inline(cell.paragraphs[0], content)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(8)
                if i == 0:
                    run.bold = True
    tr = t.rows[0]._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr.append(el)


def export(source, target):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin = sec.right_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    for key, sz in [("Title", 16), ("Heading 1", 13), ("Heading 2", 11.5), ("Heading 3", 10.5)]:
        style = doc.styles[key]
        style.font.name = "Aptos Display"
        style.font.size = Pt(sz)
        style.font.color.rgb = RGBColor(31, 53, 74)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
    header = sec.header.paragraphs[0]
    header.text = "AEROFOIL ML | AUTHOR REVIEW | 14 OCTOBER 2026"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Coauthor draft | Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line or line == "---":
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i])
                i += 1
            table(doc, rows)
            continue
        im = re.fullmatch(r"!\[(.*?)\]\(([^)]+)\)", line)
        if im:
            img = (source.parent / im.group(2)).resolve()
            if img.exists():
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.add_run().add_picture(str(img), width=Cm(16))
            i += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)", line)
        if heading:
            n = len(heading.group(1))
            key = {1: "Title", 2: "Heading 1", 3: "Heading 2"}.get(n, "Heading 3")
            inline(doc.add_paragraph(style=key), heading.group(2))
            i += 1
            continue
        if line.startswith("- "):
            inline(doc.add_paragraph(style="List Bullet"), line[2:])
            i += 1
            continue
        if line.startswith("> "):
            inline(doc.add_paragraph(style="Quote"), line[2:])
            i += 1
            continue
        block = [line]
        i += 1
        while (
            i < len(lines)
            and lines[i].strip()
            and not lines[i].startswith(("#", "|", "- ", "![", "---"))
        ):
            block.append(lines[i].strip())
            i += 1
        inline(doc.add_paragraph(), " ".join(block))
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    print(target, "paragraphs", len(doc.paragraphs), "tables", len(doc.tables))


if __name__ == "__main__":
    export(ROOT / "paper/manuscript.md", OUT / "MANUSCRIPT_AUTHOR_REVIEW.docx")
    export(OUT / "REVIEW_REPORT.md", OUT / "SCIENTIFIC_REVIEW_REPORT.docx")
