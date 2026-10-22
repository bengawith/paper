"""Final auditable source/model/academic-author-draft qualification.

Review-complete numerical study is NOT independent RANS CFD or a journal
submission decision. Persist one reproducible receipt, verify on subsequent runs.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
from pathlib import Path

from docx import Document

from robust_airfoil.modelling.freeze import verify_frozen_artifacts

ROOT = Path(__file__).resolve().parents[2]
RESULTS = ROOT / "results/robust_v2/lineages/robust-v2-controlled-20261020"
OLD = ROOT / "reports/robust_v2/external_academic_review_20261021"
LATEST = ROOT / "reports/robust_v2/final_release_20261022"
MANUSCRIPT = ROOT / "paper/manuscript.md"
WORD = ROOT / "paper/LATEST_MANUSCRIPT_FOR_AVA.docx"
SUPPLEMENT = ROOT / "paper/SUPPLEMENTARY_NUMERICAL_COMPARISON.md"
SUPP_WORD = ROOT / "paper/SUPPLEMENTARY_NUMERICAL_COMPARISON.docx"
OUTPUT = LATEST / "FINAL_REVIEW_ACCEPTANCE.json"


def sha(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1048576), b""):
            digest.update(block)
    return digest.hexdigest()


def read(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def references(text: str) -> set[int]:
    citations: set[int] = set()
    for expr in re.findall(r"\[([0-9,\s–-]+)\]", text):
        for chunk in expr.split(","):
            values = chunk.strip().replace("-", "–")
            if not values:
                continue
            if "–" in values:
                start, end = (int(x) for x in values.split("–"))
                if not 1 <= start <= end <= 29:
                    raise ValueError("Invalid bibliography reference range")
                citations.update(range(start, end + 1))
            else:
                citations.add(int(values))
    return citations


def verify() -> dict:
    source = MANUSCRIPT.read_text(encoding="utf-8")
    appendix = SUPPLEMENT.read_text(encoding="utf-8")
    historical_paper = LATEST / "PRE_FINAL_20261021_MANUSCRIPT.md"
    if not historical_paper.is_file():
        raise FileNotFoundError("Previous author manuscript must remain preserved")
    if (
        source.count("### 5.2.1 Post hoc benchmark") != 1
        or source.count("### 5.2.2 Geometry-distinct") != 1
        or "#### Direct comparison with NeuralFoil" in source
    ):
        raise ValueError("Study comparison must appear once and only once")
    if len(source.split()) < 9500:
        raise ValueError("The full manuscript is unexpectedly short")
    if (
        "independent cfd" not in source.lower()
        or "no complete nominal matched pair" not in source.lower()
    ):
        raise ValueError("Scientifically critical limitations missing from paper")
    citations = [int(x) for x in re.findall(r"(?m)^\[(\d+)\] ", source)]
    if citations != list(range(1, 30)) or references(source.split("## References")[0]) != set(
        citations
    ):
        raise ValueError("Not all 29 numbered citations are used and correctly ordered")
    ledger = list(
        csv.DictReader(
            (OLD / "REFERENCES_PRIMARY_VERIFIED_V2.csv").open(encoding="utf-8-sig", newline="")
        )
    )
    if len(ledger) != 29 or not all(x["bibliography_accurate"] == "True" for x in ledger):
        raise ValueError("Primary bibliographic identity was not verified")
    if not all(x in appendix for x in ["NeuralFoil", "17,968", "22,600", "AUROC"]):
        raise ValueError("Required adverse competitor results missing")
    manuscript = Document(WORD)
    supplement = Document(SUPP_WORD)
    if (len(manuscript.tables), len(manuscript.inline_shapes)) != (7, 8):
        raise ValueError("Word paper lost published tables or figures")
    if (len(supplement.tables), len(supplement.inline_shapes)) != (3, 0):
        raise ValueError("Word supplement's three benchmark tables are missing")
    for p in manuscript.paragraphs:
        if p.text.startswith(("Figure ", "Supplementary Figure ")) and not p.paragraph_format.keep_with_next:
            raise ValueError("Figure caption not linked to its image")
    word_receipt = read(LATEST / "FINAL_SOURCE_AND_WORD_HASHES_R2.json")
    if sha(MANUSCRIPT) != word_receipt["source_sha256"] or sha(WORD) != word_receipt["word_sha256"]:
        raise ValueError("Current Word bytes differ from the approved source")
    preview = read(LATEST / "LATEST_WORD_VISUAL_RECEIPT.json")
    if (
        preview["main"]["word_sha256"] != sha(WORD)
        or preview["supplement"]["word_sha256"] != sha(SUPP_WORD)
        or preview["main"]["figure_caption_image_groups"] != 8
        or preview["main"]["page_count"] != 22
        or preview["supplement"]["page_count"] != 2
    ):
        raise ValueError("Browser-derived page layout does not match final Word")
    for section in (preview["main"], preview["supplement"]):
        for item in section["pages"]:
            if sha(ROOT / item["image"]) != item["sha256"]:
                raise ValueError("Visual preview was changed")
        for item in section["contact_sheets"]:
            if sha(ROOT / item["path"]) != item["sha256"]:
                raise ValueError("Contact sheet was changed")
    cohort = read(RESULTS / "DATA_MANIFEST.json")
    original = read(RESULTS / "confirmation/CONFIRMATION_RESULTS.json")
    robust = read(RESULTS / "optimisation_production/PRODUCTION_SUMMARY.json")
    nominal = read(RESULTS / "nominal_control/NOMINAL_CONTROL_SUMMARY.json")
    selected = read(RESULTS / "selected_study/SELECTED_XFOIL_SUMMARY.json")
    if cohort["observed_states"] != 106931 or cohort["missing_requested_states"] != 34997:
        raise ValueError("Original controlled label accounting changed")
    if (
        original["observed_states"] != 17968
        or original["requested_states"] != 22600
        or original["claims"]["confirmation_used_to_fit_or_select_model"]
    ):
        raise ValueError("Original confirmation scientific boundaries changed")
    if (
        robust["status"] != "FULL_PROFILE_COMPLETED"
        or len(robust["per_seed"]) != 10
        or not robust["both_numerical_checks_met"]
    ):
        raise ValueError("Full original robust qualification incomplete")
    if nominal["status"] != "FULL_NOMINAL_PROFILE_COMPLETED" or len(nominal["per_seed"]) != 10:
        raise ValueError("Full nominal budget incomplete")
    if (
        selected["cases_requested"],
        selected["usable_cases"],
        selected["cases_failed_or_not_admitted"],
    ) != (819, 290, 529):
        raise ValueError("Adverse strict solver-case accounting changed")
    frozen = RESULTS / "FROZEN_MODEL_MANIFEST.json"
    manifest = read(frozen)
    verify_frozen_artifacts(frozen, {k: Path(v) for k, v in manifest["artifact_paths"].items()})
    full = read(LATEST / "NEURALFOIL_FULL_CONFIRMATION/RESULTS.json")
    group = read(LATEST / "EXTERNAL_NACA_FAMILY/NEW_FAMILY_RESULTS.json")
    group_audit = read(LATEST / "EXTERNAL_NACA_FAMILY/INDEPENDENT_RECEIPT_AND_TRUST_AUDIT.json")
    simple = read(LATEST / "NEURALFOIL_COMPARISON/BENCHMARK_RESULT.json")
    if (full["geometries_completed_by_neuralfoil"], full["same_solver_observed_paired_states"]) != (
        200,
        17968,
    ):
        raise ValueError("Full competitor diagnostic did not complete")
    if not all(
        full["per_target_observed_label_error"][target]["neuralfoil_mae"]
        < full["per_target_observed_label_error"][target]["controlled_mae"]
        for target in ("cl", "cd", "cm")
    ):
        raise ValueError("The paper's adverse comparator conclusion changed")
    if simple["paired_observed_states"] != 1013 or simple["case_geometries_attempted"] != 357:
        raise ValueError("Fixed-incidence comparison incomplete")
    if (group["requested_geometries"], group["observed_states"]) != (40, 4331) or not group_audit[
        "all_40_geometry_receipts_hash_verified"
    ]:
        raise ValueError("External 40-geometry source-bound numerical challenge incomplete")
    if not all(
        group["metrics"][f"neuralfoil_{target}_mae"] < group["metrics"][f"ours_{target}_mae"]
        for target in ("cl", "cd", "cm")
    ):
        raise ValueError("Source-distinct comparator outcome changed")
    legacy = read(OLD / "NEW_MANUSCRIPT_NUMERIC_CLAIM_AUDIT.json")
    for claim in legacy["claims"]:
        file = ROOT / claim["source_path"]
        if sha(file) != claim["source_sha256"]:
            raise ValueError("Prior 46-claim frozen source evidence changed: " + claim["claim"])
    if legacy["checks"] != 46:
        raise ValueError("Old independent source-claim audit is missing")
    if "NeuralFoil" not in source or "4,331" not in source or "22,600" not in source:
        raise ValueError("The final numerical addition was not inserted into the paper")
    paths = [
        MANUSCRIPT,
        WORD,
        SUPPLEMENT,
        SUPP_WORD,
        LATEST / "LATEST_WORD_VISUAL_RECEIPT.json",
        LATEST / "NEURALFOIL_FULL_CONFIRMATION/RESULTS.json",
        LATEST / "EXTERNAL_NACA_FAMILY/INDEPENDENT_RECEIPT_AND_TRUST_AUDIT.json",
        OLD / "REFERENCES_PRIMARY_VERIFIED_V2.csv",
    ]
    return {
        "schema": "robust-v2-final-coauthor-review-gate-v1",
        "status": "READY_FOR_COAUTHOR_REVIEW_INDEPENDENT_CFD_PENDING",
        "numeric_study": "COMPLETE_NO_RETRAIN_OR_SEED_RERUN",
        "original_heldout_confirmed_observed_states": 17968,
        "independent_naca_source_geometries": 40,
        "new_family_observed_numeric_states": 4331,
        "strict_selected_xfoil_failed_or_unusable": 529,
        "benchmark_neuralfoil_advantage_honestly_reported": True,
        "reference_primary_identity_count": 29,
        "original_science_reconciled_claim_count": 46,
        "main_manuscript_words": len(source.split()),
        "main_word_tables": len(manuscript.tables),
        "main_word_figures": len(manuscript.inline_shapes),
        "supplement_tables": len(supplement.tables),
        "browser_preview_main_pages": 22,
        "browser_preview_supplement_pages": 2,
        "browser_derived_not_native_word": True,
        "journal_submission_ready": False,
        "independent_cfd_performed": False,
        "author_contributions_approved": False,
        "files_sha256": {path.relative_to(ROOT).as_posix(): sha(path) for path in paths},
        "remaining": "Independent CFD results, CFD-dependent final physical conclusion, and coauthor disclosures/native Word check",
    }


def main() -> None:
    result = verify()
    if OUTPUT.exists():
        if read(OUTPUT) != result:
            raise ValueError("Existing review receipt differs from recomputed evidence")
    else:
        OUTPUT.write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in result.items() if k != "files_sha256"}, indent=2))


if __name__ == "__main__":
    main()
