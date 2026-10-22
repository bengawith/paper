"""Fix original author prose as an immutable style source, not a result authority."""

import hashlib
import json
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent
ORIGINAL = Path("C:/Users/Ben/OneDrive/paper/prior_final_draft.docx")
PAPER = ROOT / "paper/manuscript.md"
WORD = ROOT / "paper/LATEST_MANUSCRIPT_FOR_AVA.docx"


def h(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


for name, source in [
    ("PRE_REWRITE_CONTROLLED_MANUSCRIPT.md", PAPER),
    ("PRE_REWRITE_CONTROLLED_WORD.docx", WORD),
    ("ORIGINAL_USER_SUBMISSION.docx", ORIGINAL),
]:
    target = OUT / name
    if target.exists():
        raise FileExistsError("Preserve existing baseline: " + str(target))
    target.write_bytes(source.read_bytes())
doc = Document(ORIGINAL)
titles = [(i, p.text) for i, p in enumerate(doc.paragraphs) if p.style.name.startswith("Heading")]
sections = {}
for index, (i, title) in enumerate(titles):
    end = titles[index + 1][0] if index + 1 < len(titles) else len(doc.paragraphs)
    paragraphs = [
        p.text for p in doc.paragraphs[i + 1 : end] if p.text.strip() and p.style.name == "Normal"
    ]
    sections[title] = dict(
        paragraphs=len(paragraphs),
        samples=paragraphs[:3],
        word_count=sum(len(v.split()) for v in paragraphs),
    )
style = {
    "source": "C:/Users/Ben/OneDrive/paper/prior_final_draft.docx",
    "source_sha256": h(ORIGINAL),
    "original_paragraphs": len(doc.paragraphs),
    "original_tables": len(doc.tables),
    "original_inline_figures": len(doc.inline_shapes),
    "previous_controlled_source_sha256": h(PAPER),
    "previous_controlled_word_sha256": h(WORD),
    "stylistic_rules": [
        "Use first-person plural in methodological choices and report the actual scope.",
        "Introduce the engineering problem, describe a short evidence-based gap, state explicit contributions.",
        "Define data sources, geometry and targets before architecture and numerical results.",
        "Explain why a methodological choice is needed before the detailed algorithm.",
        "Use descriptive subsection headings, physical context and compact quantitative claims.",
        "Use contributions and explanations from the original voice; do not retain its unsupported numerical claims, repetitions or ambiguous symbol equations.",
        "Do not invent coauthor writing styles, declarations or physical CFD evidence.",
    ],
    "sections": sections,
}
(OUT / "ORIGINAL_AUTHOR_STYLE_EVIDENCE.json").write_text(
    json.dumps(style, ensure_ascii=False, indent=2), encoding="utf-8"
)
print(
    "AUTHOR_STYLE_SOURCE_SHA",
    h(ORIGINAL),
    "PARAGRAPHS",
    len(doc.paragraphs),
    "SECTIONS",
    len(sections),
    flush=True,
)
