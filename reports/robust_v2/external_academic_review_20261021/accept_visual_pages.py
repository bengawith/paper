"""Document 21-page inspected nonnative Word-content visual QA."""

import hashlib
import json
from pathlib import Path

import pypdfium2
from docx import Document

ROOT = Path(__file__).resolve().parents[3]
HERE = Path(__file__).resolve().parent
WORD = ROOT / "paper/LATEST_MANUSCRIPT_FOR_AVA.docx"
PDF = HERE / "RENDERED_V2/STANDALONE_AUTHOR_VOICE_FINAL.browser_preview.pdf"
OUT = HERE / "VISUAL_QA_FULL_AUTHOR_STYLE.json"


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


if OUT.exists():
    raise FileExistsError(OUT)
doc = Document(WORD)
pages = len(pypdfium2.PdfDocument(str(PDF)))
assert len(doc.tables) == 7 and len(doc.inline_shapes) == 8 and pages == 21
files = [HERE / f"RENDERED_V2/PAGE_{i:02d}.png" for i in range(1, 22)]
sheets = [HERE / f"RENDERED_V2/CONTACT_SHEET_{i:02d}.png" for i in range(1, 5)]
assert all(p.is_file() for p in files + sheets)
a = {
    "schema": "full-rewrite-browser-page-visual-audit-v1",
    "status": "BROWSER_DERIVED_PAGES_REVIEWED",
    "pages": pages,
    "word_figures": 8,
    "word_tables": 7,
    "captions_grouped_with_images": True,
    "manuscript_sha256": sha(ROOT / "paper/manuscript.md"),
    "word_sha256": sha(WORD),
    "pdf_sha256": sha(PDF),
    "all_21_pages_visually_inspected": True,
    "four_contact_sheets_visually_inspected": True,
    "inspection_note": "All figures with captions on the same preview page, readable tables, no visibly clipped text; checked all page/contact-sheet previews.",
    "native_microsoft_word_pagination_verified": False,
    "limitations": "Chromium derived from true DOCX contents and image bytes; Microsoft Word pagination may differ.",
    "page_sha256": {p.name: sha(p) for p in files},
    "contact_sheet_sha256": {p.name: sha(p) for p in sheets},
}
OUT.write_text(json.dumps(a, indent=2), encoding="utf-8")
print("ALL_21_PAGE_PREVIEW_QA_PASS; NATIVE_WORD_UNVERIFIED")
