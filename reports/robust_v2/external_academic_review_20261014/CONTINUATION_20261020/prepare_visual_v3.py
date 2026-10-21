"""Prepare per-page visual quality evidence for the exact exported Word documents.

Render a DOCX-content-derived Chromium preview with all embedded figures,
then inspect PNG page renders. This is NOT native Microsoft Word paginated
layout; that limitation must remain explicit in the review.
"""

from __future__ import annotations

import hashlib
import json
import subprocess
import sys
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document
from PIL import Image, ImageDraw, ImageOps

ROOT = Path(__file__).resolve().parents[4]
REVIEW = ROOT / "reports/robust_v2/external_academic_review_20261014"
HERE = Path(__file__).resolve().parent
RENDERED = REVIEW / "RENDERED_V3"
SOURCE = HERE / "render_author_docx_browser.py"


def sha(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1048576), b""):
            h.update(chunk)
    return h.hexdigest()


def document(name, word):
    if not word.is_file():
        raise FileNotFoundError("Word export not yet complete: " + str(word))
    doc = Document(word)
    pdf = RENDERED / f"{name}.browser-preview.pdf"
    if pdf.exists():
        raise FileExistsError("Existing preview must be inspected: " + str(pdf))
    subprocess.run(
        [sys.executable, str(SOURCE), str(word), "--output", str(pdf)], cwd=ROOT, check=True
    )
    reader = pdfium.PdfDocument(str(pdf))
    images = []
    pages = []
    for i, page in enumerate(reader):
        pil = page.render(scale=1.3).to_pil().convert("RGB")
        result = RENDERED / f"{name}-page-{i + 1:02d}.png"
        pil.save(result, optimize=True)
        gray = ImageOps.grayscale(pil)
        empty = sum(1 for v in gray.resize((80, 114)).getdata() if v > 247) / 9120
        pages.append(
            {
                "page": i + 1,
                "image": result.relative_to(ROOT).as_posix(),
                "preview_white_fraction_approx": round(empty, 4),
                "width_px": pil.width,
                "height_px": pil.height,
                "visually_inspected": False,
            }
        )
        thumb = pil.copy()
        thumb.thumbnail((500, 720))
        images.append(thumb)
    for offset in range(0, len(images), 6):
        sample = images[offset : offset + 6]
        sheet = Image.new("RGB", (1540, 1570), "#e4e7eb")
        draw = ImageDraw.Draw(sheet)
        for j, pil in enumerate(sample):
            col, row = j % 3, j // 3
            x = 10 + col * 510 + (500 - pil.width) // 2
            y = 30 + row * 770
            sheet.paste(pil, (x, y))
            draw.text(
                (15 + col * 510, 10 + row * 770), f"{name}: page {offset + j + 1}", fill="black"
            )
        sheet.save(RENDERED / f"{name}-contact-sheet-{offset // 6 + 1:02d}.png", optimize=True)
    return {
        "docx": word.relative_to(ROOT).as_posix(),
        "docx_sha256": sha(word),
        "structural_tables": len(doc.tables),
        "embedded_figures": len(doc.inline_shapes),
        "preview_pdf": pdf.relative_to(ROOT).as_posix(),
        "preview_pdf_sha256": sha(pdf),
        "page_count": len(reader),
        "pages": pages,
        "native_word_render": False,
        "preview_type": "Chromium HTML reconstruction from actual DOCX paragraphs,tables,images",
        "limitation": "Does not prove native Microsoft Word page breaks, headers, footers or exact typography.",
    }


def main():
    if RENDERED.exists() and any(RENDERED.iterdir()):
        raise FileExistsError(
            "Existing rendered review pages must be preserved and audited before generating again"
        )
    # The review directory may be created by an earlier read-only visual audit.
    # An existing EMPTY directory is safe; existing files above still fail closed.
    RENDERED.mkdir(parents=True, exist_ok=True)
    manuscript = document("manuscript", ROOT / "paper/LATEST_MANUSCRIPT_FOR_AVA.docx")
    referee = document("referee", REVIEW / "SCIENTIFIC_REVIEW_REPORT.docx")
    if manuscript["embedded_figures"] != 8 or manuscript["structural_tables"] != 7:
        raise ValueError("Current Word figure or table structure incomplete")
    status = {
        "schema": "word-content-visual-preview-v1",
        "manuscript_docx_sha256": manuscript["docx_sha256"],
        "referee_docx_sha256": referee["docx_sha256"],
        "visual_previews_created": True,
        "all_pages_visually_inspected": False,
        "native_microsoft_word_render_unavailable": True,
        "manuscript": manuscript,
        "referee": referee,
        "reviewer_action": "Inspect every generated page PNG and contact sheet, record abnormalities and re-render. Do not claim native Word layout parity.",
    }
    (REVIEW / "VISUAL_QA_V3.json").write_text(json.dumps(status, indent=2), encoding="utf-8")
    print(
        "WORD_CONTENT_PREVIEWS_READY_FOR_MANUAL_INSPECTION",
        manuscript["page_count"],
        referee["page_count"],
        flush=True,
    )


if __name__ == "__main__":
    main()
