"""Independent Chromium layout preview from exact Word DOCX contents.

This is a DOCX-derived layout preview, not Microsoft Word's native rendering.
Only mark visual QA after manually inspecting all generated PDF pages and the
actual Word structure, and label the renderer limitation accurately.
"""

from __future__ import annotations

import base64
import html
import json
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[4]
HERE = Path(__file__).resolve().parent
CHROME = Path("C:/Program Files/Google/Chrome/Application/chrome.exe")


def inline(paragraph, doc):
    images = []
    for blip in paragraph._p.xpath(".//a:blip"):
        ident = blip.get(qn("r:embed"))
        if ident:
            part = doc.part.related_parts[ident]
            image = (
                '<img class="figure" alt="Embedded Word figure" src="data:'
                + part.content_type
                + ";base64,"
                + base64.b64encode(part.blob).decode("ascii")
                + '">'
            )
            images.append(image)
    text = []
    for run in paragraph.runs:
        span = html.escape(run.text).replace("\n", "<br>")
        if run.bold:
            span = "<strong>" + span + "</strong>"
        if run.italic:
            span = "<em>" + span + "</em>"
        if run.font.superscript:
            span = "<sup>" + span + "</sup>"
        if run.font.subscript:
            span = "<sub>" + span + "</sub>"
        text.append(span)
    return "".join(text) + "".join(images)


def html_doc(path):
    doc = Document(path)
    body = []
    image_count = 0
    for element in doc.element.body:
        if element.tag == qn("w:p"):
            par = Paragraph(element, doc)
            name = par.style.name or "Normal"
            content = inline(par, doc)
            if "<img" in content:
                image_count += content.count("<img")
            if name == "Title":
                body.append("<h1>" + content + "</h1>")
            elif name.startswith("Heading"):
                level = 2 if name == "Heading 1" else 3 if name == "Heading 2" else 4
                body.append(f"<h{level}>{content}</h{level}>")
            elif name == "List Bullet":
                body.append('<p class="bullet">• ' + content + "</p>")
            else:
                body.append("<p>" + content + "</p>")
        elif element.tag == qn("w:tbl"):
            table = Table(element, doc)
            rows = []
            for ix, row in enumerate(table.rows):
                kind = "th" if ix == 0 else "td"
                rows.append(
                    "<tr>"
                    + "".join(
                        f"<{kind}>"
                        + "".join(inline(p, doc) for p in cell.paragraphs)
                        + f"</{kind}>"
                        for cell in row.cells
                    )
                    + "</tr>"
                )
            body.append(
                "<table><thead>"
                + rows[0]
                + "</thead><tbody>"
                + "".join(rows[1:])
                + "</tbody></table>"
            )
    css = """@page{size:A4;margin:18mm 19mm 20mm 19mm}
 html,body{font:10.5pt/1.3 'Aptos','Calibri','Arial',sans-serif;color:#172b3c}
 body{margin:0}h1{font-size:18pt;line-height:1.16;margin:0 0 11pt;color:#18304a;break-after:avoid}
 h2{font-size:13.5pt;margin:14pt 0 7pt;color:#18304a;break-after:avoid}
 h3{font-size:11.5pt;margin:10pt 0 6pt;color:#18304a;break-after:avoid}
 h4{font-size:11pt;margin:10pt 0 5pt;color:#18304a;break-after:avoid}
 p{margin:0 0 6pt;orphans:3;widows:3}p.bullet{padding-left:11pt;text-indent:-10pt}
 table{border-collapse:collapse;width:100%;font-size:9pt;margin:7pt 0 10pt;break-inside:auto}
 thead{display:table-header-group}tr{break-inside:avoid}td,th{border:1px solid #b6c2cc;padding:5px 6px;text-align:left;vertical-align:top}
 th{background:#edf2f5}img.figure{display:block;max-width:100%;max-height:205mm;width:auto;height:auto;margin:8pt auto;object-fit:contain;break-inside:avoid}
 a{color:#183b5b}strong{font-weight:700}
 """
    return (
        '<!doctype html><html><head><meta charset="utf-8"><style>'
        + css
        + "</style></head><body>"
        + "".join(body)
        + "</body></html>",
        image_count,
        len(doc.tables),
        len(doc.paragraphs),
    )


def main():
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("word", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    word = args.word.resolve()
    output = (args.output or word.with_suffix(".browser-preview.pdf")).resolve()
    if not word.is_file() or not CHROME.is_file():
        raise FileNotFoundError("Word source or Chromium renderer unavailable")
    if output.exists():
        raise FileExistsError("A preview already exists; preserve and inspect it")
    markup, images, tables, paragraphs = html_doc(word)
    with tempfile.TemporaryDirectory(prefix="paper-visual-browser-") as path:
        folder = Path(path)
        source = folder / "document.html"
        source.write_text(markup, encoding="utf-8")
        profile = folder / "chrome-profile"
        args = [
            str(CHROME),
            "--headless=new",
            "--disable-gpu",
            "--no-first-run",
            "--disable-extensions",
            "--no-default-browser-check",
            "--no-pdf-header-footer",
            "--user-data-dir=" + str(profile),
            "--print-to-pdf=" + str(output),
            source.as_uri(),
        ]
        completed = subprocess.run(args, capture_output=True, text=True, timeout=120, check=False)
        if completed.returncode or not output.is_file() or output.stat().st_size < 1000:
            raise RuntimeError(
                "Chromium preview failed: " + completed.stdout[-1500:] + completed.stderr[-1500:]
            )
        print(
            json.dumps(
                {
                    "status": "DOCX_DERIVED_BROWSER_LAYOUT_PREVIEW_ONLY",
                    "source": str(word),
                    "preview": str(output),
                    "figures": images,
                    "tables": tables,
                    "paragraphs": paragraphs,
                    "bytes": output.stat().st_size,
                    "not_native_microsoft_word_render": True,
                },
                indent=2,
            )
        )


if __name__ == "__main__":
    main()
