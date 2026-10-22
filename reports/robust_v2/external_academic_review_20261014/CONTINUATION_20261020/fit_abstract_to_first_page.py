"""Final concise academic abstract preserving the source-bound numbers and limits.

Prevents isolated abstract tail in browser Word-content preview. Does not
change any scientific result, model, source, tests, or simulated conditions.
"""
from pathlib import Path
import hashlib,json,runpy
from docx import Document
ROOT=Path(__file__).resolve().parents[4]
HERE=Path(__file__).resolve().parent
REVIEW=ROOT/'reports/robust_v2/external_academic_review_20261014'
PAPER=ROOT/'paper'
def sha(p):
 h=hashlib.sha256()
 with p.open('rb') as f:
  for block in iter(lambda:f.read(1048576),b''):h.update(block)
 return h.hexdigest()
def main():
 p=PAPER/'manuscript.md'
 word=PAPER/'LATEST_MANUSCRIPT_FOR_AVA.docx'
 backup=HERE/'BEFORE_ABSTRACT_FIT_v3.md'
 docbackup=HERE/'BEFORE_ABSTRACT_FIT_v3.docx'
 evidence=REVIEW/'EVIDENCE_SHA256_MANIFEST.json'
 evidence_backup=HERE/'BEFORE_ABSTRACT_FIT_EVIDENCE_v3.json'
 receipt=HERE/'ABSTRACT_FIT_V3_RECEIPT.json'
 if any(x.exists() for x in [backup,docbackup,evidence_backup,receipt]):
  raise FileExistsError('A third copy-edit revision already exists')
 old=p.read_text(encoding='utf-8')
 changes=[
 ('Aerofoil surrogate models can interpolate numerical flow labels accurately without necessarily supporting reliable geometry selection.',
  'Accurate surrogate interpolation does not establish reliable aerofoil selection.'),
 ('Additional solver comparisons show that reorienting a source aerofoil can change the aerodynamic response at the same recorded angle of attack.',
  'Changing coordinate orientation can alter the XFOIL response at fixed recorded incidence.'),
 ('Instead of reusing those uncertain geometry–label pairings, we assemble a controlled, fixed-operating-condition dataset using local XFOIL calculations performed on the exact twelve-parameter Class–Shape Transformation reconstructions supplied to the model.',
  'We therefore generate fresh, fixed-condition XFOIL labels on the exact twelve-parameter CST reconstructions supplied to the model.'),
 ('Because the labels and verification share a numerical solver family, and the prescribed manufacturing field is not measured, these outcomes are conditional numerical evidence, not independent physical validation or proof of global aerodynamic optimality.',
  'Same-family solver labels and unmeasured manufacturing variation limit these to conditional numerical findings, not independent physical validation or global optimality.'),
 ]
 updated=old
 for before,after in changes:
  if updated.count(before)!=1:raise ValueError('Unexpected abstract wording, preserve paper')
  updated=updated.replace(before,after)
 if old.split('## Nomenclature')[1]!=updated.split('## Nomenclature')[1]:
  raise AssertionError('Copy edit touched manuscript beyond abstract')
 backup.write_bytes(p.read_bytes())
 docbackup.write_bytes(word.read_bytes())
 evidence_backup.write_bytes(evidence.read_bytes())
 p.write_text(updated,encoding='utf-8')
 exporter=runpy.run_path(str(REVIEW/'export_current_review.py'))['export']
 exporter(p,word)
 d=Document(word)
 if len(d.tables)!=7 or len(d.inline_shapes)!=8:raise ValueError('Word science content structure changed')
 source=json.loads(evidence.read_text(encoding='utf-8'))
 for entry in source['files']:
  if entry['relative_path'] in {'paper/manuscript.md','paper/LATEST_MANUSCRIPT_FOR_AVA.docx'}:
   f=ROOT/entry['relative_path']
   entry['sha256']=sha(f)
   entry['bytes']=f.stat().st_size
 evidence.write_text(json.dumps(source,indent=2)+'\n',encoding='utf-8')
 receipt.write_text(json.dumps({'schema':'abstract-clarity-layout-fit-v3',
  'old_paper_sha256':sha(backup),'new_paper_sha256':sha(p),
  'old_docx_sha256':sha(docbackup),'new_docx_sha256':sha(word),
  'old_manifest_sha256':sha(evidence_backup),'new_manifest_sha256':sha(evidence),
  'scientific_quantities_changed':False,'frozen_simulation_data_changed':False,
  'word_figure_count':len(d.inline_shapes),'word_table_count':len(d.tables),
  'visual_qa_required':True},indent=2),encoding='utf-8')
 print('ABSTRACT_LAYOUT_FIT_V3_READY',sha(word))
if __name__=='__main__':main()
