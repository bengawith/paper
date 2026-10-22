"""Reconcile the current referee report, model card, Word export and evidence.

Requires terminal full campaign/confirmation/selected receipts. Never claims
independent CFD, external peer review, or coauthor approval.
"""

from __future__ import annotations

import csv
import hashlib
import json
import runpy
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[4]
BASE = ROOT / "reports/robust_v2/external_academic_review_20261014"
CONTROL = BASE / "CONTROLLED_MODEL"
LINEAGE = ROOT / "results/robust_v2/lineages/robust-v2-controlled-20261020"
PAPER = ROOT / "paper"


def load(path):
    return json.loads(path.read_text(encoding="utf-8"))


def digest(p):
    h = hashlib.sha256()
    with p.open("rb") as f:
        for block in iter(lambda: f.read(1048576), b""):
            h.update(block)
    return h.hexdigest()


def as_text(path, text):
    path.write_text(text.strip() + "\n", encoding="utf-8")


def number(n):
    return "not estimable" if n is None else f"{float(n):.4f}"


def main():
    data = load(LINEAGE / "DATA_MANIFEST.json")
    test = load(CONTROL / "REPAIRED_TEST_METRICS.json")
    risk = load(CONTROL / "RISK_OOF/REPAIRED_RISK_OOF.json")
    confirmation = load(CONTROL / "CONFIRMATION/CONFIRMATION_RESULTS.json")
    production = load(LINEAGE / "optimisation_production/PRODUCTION_SUMMARY.json")
    nominal = load(LINEAGE / "nominal_control/NOMINAL_CONTROL_SUMMARY.json")
    selected = load(LINEAGE / "selected_study/SELECTED_XFOIL_SUMMARY.json")
    if (
        production.get("status") != "FULL_PROFILE_COMPLETED"
        or len(production.get("per_seed", [])) != 10
    ):
        raise ValueError("Ten full robust seeds are required")
    if (
        nominal.get("status") != "FULL_NOMINAL_PROFILE_COMPLETED"
        or len(nominal.get("per_seed", [])) != 10
    ):
        raise ValueError("Ten nominal-control seeds are required")
    if confirmation["requested_geometries"] != 200:
        raise ValueError("Sealed confirmation is incomplete")
    if (
        selected["cases_requested"]
        != selected["usable_cases"] + selected["cases_failed_or_not_admitted"]
    ):
        raise ValueError("Missing or duplicated selected-XFOIL outcomes")
    manuscript = PAPER / "manuscript.md"
    text = manuscript.read_text(encoding="utf-8")
    if (
        "{r2_cl}" in text
        or "Controlled exact-CST" not in text
        and "Geometric Provenance" not in text
    ):
        raise ValueError("Manuscript not yet reconciled to terminal controlled evidence")
    achieved = production["both_numerical_checks_met"]
    decision = (
        "both fixed multi-seed numerical comparisons pass"
        if achieved
        else "the fixed multi-seed numerical comparison is not fully satisfied"
    )
    referee = f"""# Independent computational academic-integrity review — final controlled study

21 October 2026. Editorial recommendation: **MAJOR REVISIONS BEFORE SUBMISSION**.

This is a separate computational/source-grounded self-review and not a fabricated independent human peer referee or coauthor endorsement. This manuscript is an author-review draft, not a journal acceptance.

## Source defects, scope and positive evidence

The earlier coordinate-order defect affected 1,001/1,560 encoded aerofoils. The subsequent exact-geometry numerical corpus contains **{data["observed_states"]:,} actual local-XFOIL numerical labels** from {data["geometries_with_observations"]:,} CST geometries, out of {data["requested_unique_states"]:,} requested states. Missing or unusable cases total {data["missing_requested_states"]:,}; only {data["complete_grid_geometries"]:,} shape sweeps are complete. The omitted states are not estimated, filled or treated as random.

Five fixed-architecture neural members were retrained on genuine local solver labels, stopped using only out-of-bag development clusters, and hash-frozen before protected test evaluation. The held-out historical-identity numerical test contains {test["states"]:,} observed states from {test["airfoils"]:,} aerofoils. These identities were previously inspected and are not a pristine independent family sample.

On the controlled calibration data, support-only AUROC is {number(risk["models"]["support_only"]["auroc"])}; combined support-plus-disagreement AUROC is {number(risk["models"]["support_plus_ensemble"]["auroc"])}. The combined score contains ranking information, not a newly deployed guaranteed coverage policy.

A separately sealed 200-geometry confirmation cohort, chosen before training and predicted before its labels, yielded {confirmation["observed_states"]:,}/{confirmation["requested_states"]:,} clean numerical outcomes. The frozen incremental high-error ranking AUROC is {number(confirmation["incremental_combined_minus_support_auroc"])}. The ensemble-disagreement-only AUROC is {number(confirmation["risk_models"]["ensemble_only"]["auroc"])}, compared with {number(confirmation["risk_models"]["support_plus_ensemble"]["auroc"])} for the combined model: the combined signal is not necessarily the strongest risk ranking. The cohort consists of source-development geometry mixtures, not independent aircraft families. All solver failures remain in the denominator of requested cases.

The full candidate-specific NSGA-II profile ran ten seeds at population128/generations200/scenarios64; a matched constrained nominal-only ten-seed control also completed at 128×200. The predeclared reporting verdict is that {decision}. Reference-scaled distance is {number(production["reference_scaled_agreement"]["pairwise_objective_distance_mean"])}; historical pooled-range distance is {number(production["original_pooled_range_diagnostic"]["pairwise_objective_distance_mean"])}. A finite seeded search, successful or unsuccessful, cannot establish global optimality.

Fresh strict selected-design XFOIL was attempted for {selected["cases_requested"]:,} cases. {selected["usable_cases"]:,} had all fixed required-lift and solver-completion conditions; {selected["cases_failed_or_not_admitted"]:,} did not. The final paper keeps paired-success denominators, adverse tails and any unfavorable ranking reversals.

## Independent research readiness decision

The software reconstruction and numerical evidence are auditable. **Independent viscous RANS CFD and physical experiments have not been performed.** A measured manufacturing distribution is unavailable; the field remains prescribed sensitivity. Ncrit is not itself a RANS transition/turbulence specification. No journal submission, coauthor CRediT confirmation, funding/conflict disclosure approval or probability of acceptance is claimed.

Ava and the authors must confirm the physical mission and boundary conditions, run converged mesh/domain/transition-sensitive independent CFD at matched lift, retain full force and moment histories and failed cases, examine any changed ranking, approve declarations and choose the target journal. This remains a major-revision scientific gate, not something a code/test/packaging pass can waive.

## Traceability and artefact acceptance

Source fixes, training hashes, frozen cohort identity, before-label risk scores, the exact numerical loss/failure accounting, ten seeded robust and nominal records, 2,048-scenario sensitivity and fresh selected-case receipts are retained in the repository/evidence archives. The current paper and report are separate standalone Word files; source/claim matrices and the SHA256 manifest are the authority. The authors should inspect both in native Word, including every page, before distributing.

**Editorial recommendation: MAJOR REVISIONS BEFORE SUBMISSION.**
"""
    as_text(BASE / "ACADEMIC_REFEREE_REPORT.md", referee)
    card = f"""# Controlled numerical model card

Lineage robust-v2-controlled-20261020. Five point-conditioned residual MLPs map twelve CST parameters and incidence to CL, log CD and CM. True local XFOIL labels on exact reconstructed input geometries were used for training. Source data has {data["observed_states"]:,} observed and {data["missing_requested_states"]:,} missing/invalid states. All member weights and scalers, split assignments and calibration outputs are SHA256-frozen.

This model is fit for cautious reproducible XFOIL-family numerical prediction and constrained design research only. Its held-out historical identities have been inspected, and the sealed confirmation uses source-shape mixtures. The support/ensemble score does not certify conditional/adaptive coverage. This model has not been independently validated by RANS CFD, manufacturing experiments or flight tests.
"""
    as_text(PAPER / "MODEL_CARD.md", card)
    handoff = """# Robust V2 — author review handoff

The current standalone Word manuscript is paper/LATEST_MANUSCRIPT_FOR_AVA.docx. The scientific referee report is ACADEMIC_REFEREE_REPORT.md. All five controlled models, the 200-geometry numerical confirmation, ten robust seeds, ten nominal seeds and all strict selected XFOIL cases have terminal receipts. Actual retained adverse/missing outcomes and uncertainty caveats appear in the manuscript and evidence ledger.

**Major scientific revision remains necessary before journal submission:** independent CFD, measured manufacturing interpretation and author-approved declarations/mission. Word page rendering/visual verification, final static/test gates, ordinary review-branch push and a post-push OneDrive SHA256 package are separate final delivery steps, not inferred from script generation.
"""
    as_text(BASE / "READ_THIS_FIRST.md", handoff)
    as_text(ROOT / "README.md", handoff)
    as_text(PAPER / "README.md", handoff)
    matrix = BASE / "CLAIM_EVIDENCE_MATRIX.csv"
    with matrix.open(encoding="utf-8-sig", newline="") as f:
        prior = list(csv.DictReader(f))
    if any(x.get("claim_id") == "C33" for x in prior):
        raise ValueError("Do not duplicate controlled claim rows")
    fields = list(prior[0])
    backup = BASE / "CONTINUATION_20261020" / "CLAIM_MATRIX_PRE_CONTROLLED.csv"
    if backup.exists():
        raise FileExistsError("Historical claim matrix backup already exists; do not overwrite it")
    backup.write_bytes(matrix.read_bytes())
    for claim in prior:
        path = ROOT / claim["source_path"]
        if not path.is_file():
            raise FileNotFoundError("Historical claim evidence file missing: " + str(path))
        current = digest(path)
        if current != claim["source_sha256"]:
            original_hash = claim["source_sha256"]
            claim["source_sha256"] = current
            claim["status"] = "HISTORICAL_SOURCE_SUPERSEDED_NOT_CURRENT_SCIENTIFIC_VALIDATION"
            claim["limit_or_action"] = (
                claim["limit_or_action"]
                + "; historical source SHA256 "
                + original_hash
                + "; live source revised after earlier audit; original historical numbers are not reattributed"
            )
    next_claims = [
        (
            "C33",
            "2.2;Table1",
            "Exact labelled data and missingness",
            LINEAGE / "DATA_MANIFEST.json",
            f"{data['observed_states']} of {data['requested_unique_states']} observed",
        ),
        (
            "C34",
            "4.1;Table3",
            "Controlled protected-test performance",
            CONTROL / "REPAIRED_TEST_METRICS.json",
            str(test["states"]) + " states",
        ),
        (
            "C35",
            "4.2;Table4",
            "Calibration geometry OOF risk",
            CONTROL / "RISK_OOF/REPAIRED_RISK_OOF.json",
            number(risk["incremental_auroc"]),
        ),
        (
            "C36",
            "3.3;4.2",
            "Before-label frozen confirmation",
            CONTROL / "CONFIRMATION/CONFIRMATION_RESULTS.json",
            str(confirmation["observed_states"]) + " observed",
        ),
        (
            "C37",
            "4.3;Figure4",
            "Full ten-seed risk optimization",
            LINEAGE / "optimisation_production/PRODUCTION_SUMMARY.json",
            number(production["reference_scaled_agreement"]["pairwise_objective_distance_mean"]),
        ),
        (
            "C38",
            "4.3;Table5",
            "Full ten-seed nominal control",
            LINEAGE / "nominal_control/NOMINAL_CONTROL_SUMMARY.json",
            "128x200x10",
        ),
        (
            "C39",
            "4.4;Figure7",
            "Fresh 2048-scenario sensitivity",
            LINEAGE / "selected_study/FINAL_2048_SCENARIO_EVALUATION.json",
            "2048 common scenarios",
        ),
        (
            "C40",
            "4.4;Table6",
            "Selected local solver validation",
            LINEAGE / "selected_study/SELECTED_XFOIL_SUMMARY.json",
            str(selected["usable_cases"]) + " usable of " + str(selected["cases_requested"]),
        ),
    ]
    for ident, loc, claim, path, value in next_claims:
        prior.append(
            {
                "claim_id": ident,
                "manuscript_location": loc,
                "claim": claim,
                "source_path": path.relative_to(ROOT).as_posix(),
                "source_sha256": digest(path),
                "selector": "terminal machine report",
                "recomputed_value": value,
                "status": "VERIFIED_NUMERICALLY_QUALIFIED",
                "limit_or_action": "No independent CFD; failed case accounting preserved",
            }
        )
    with matrix.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(prior)
    # Export the current standalone source; no referee text is appended to paper.
    exporter = runpy.run_path(str(BASE / "export_current_review.py"))["export"]
    exporter(manuscript, PAPER / "LATEST_MANUSCRIPT_FOR_AVA.docx")
    exporter(BASE / "ACADEMIC_REFEREE_REPORT.md", BASE / "SCIENTIFIC_REVIEW_REPORT.docx")
    doc = Document(PAPER / "LATEST_MANUSCRIPT_FOR_AVA.docx")
    if len(doc.tables) != 7 or len(doc.inline_shapes) != 8:
        raise ValueError("Word export does not contain seven tables and eight embedded figures")
    evidence = []
    for p in [
        manuscript,
        PAPER / "LATEST_MANUSCRIPT_FOR_AVA.docx",
        PAPER / "MODEL_CARD.md",
        BASE / "ACADEMIC_REFEREE_REPORT.md",
        BASE / "SCIENTIFIC_REVIEW_REPORT.docx",
        matrix,
        BASE / "REFERENCES_VERIFIED.csv",
        LINEAGE / "DATA_MANIFEST.json",
        LINEAGE / "FROZEN_MODEL_MANIFEST.json",
        CONTROL / "REPAIRED_TEST_METRICS.json",
        CONTROL / "RISK_OOF/REPAIRED_RISK_OOF.json",
        CONTROL / "CONFIRMATION/CONFIRMATION_RESULTS.json",
        LINEAGE / "optimisation_production/PRODUCTION_SUMMARY.json",
        LINEAGE / "nominal_control/NOMINAL_CONTROL_SUMMARY.json",
        LINEAGE / "selected_study/SELECTED_XFOIL_SUMMARY.json",
        BASE / "CONTINUATION_20261020/EXACT_RUNTIME_SOURCE_BYTES.zip",
        BASE / "CONTINUATION_20261020/EXACT_RUNTIME_SOURCE_MANIFEST.json",
        *sorted((PAPER / "figures").glob("fig_*.png")),
    ]:
        if not p.is_file():
            raise FileNotFoundError(str(p))
        evidence.append(
            {
                "relative_path": p.relative_to(ROOT).as_posix(),
                "sha256": digest(p),
                "bytes": p.stat().st_size,
            }
        )
    as_text(
        BASE / "EVIDENCE_SHA256_MANIFEST.json",
        json.dumps({"schema": "controlled-manuscript-evidence-v1", "files": evidence}, indent=2),
    )
    as_text(
        BASE / "FINAL_ACCEPTANCE.json",
        json.dumps(
            {
                "status": "REVIEW_COMPUTATION_COMPLETE_MAJOR_REVISIONS_REQUIRED",
                "review_deliverables_complete": False,
                "software_checks": "PENDING_FINAL_STATIC_TESTS_AND_WORD_VISUAL_QA",
                "controlled_model_training": "COMPLETE",
                "sealed_confirmation": "COMPLETE",
                "full_ten_seed_robust": "COMPLETE",
                "full_ten_seed_nominal": "COMPLETE",
                "selected_cases_requested": selected["cases_requested"],
                "selected_cases_usable": selected["usable_cases"],
                "independent_cfd": "NOT_PERFORMED",
                "journal_submission_ready": False,
                "coauthor_approval": "NOT_CLAIMED",
                "manuscript_sha256": digest(manuscript),
                "docx_sha256": digest(PAPER / "LATEST_MANUSCRIPT_FOR_AVA.docx"),
                "embedded_figures": len(doc.inline_shapes),
                "tables": len(doc.tables),
            },
            indent=2,
        ),
    )
    print(
        "FINAL_REFEREE_WORD_EVIDENCE_DRAFT_CREATED", len(evidence), "hashed artifacts", flush=True
    )


if __name__ == "__main__":
    main()
