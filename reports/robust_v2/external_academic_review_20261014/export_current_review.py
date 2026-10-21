"""Export only the current standalone manuscript and its separate referee report."""
from pathlib import Path
import hashlib
import json
import runpy
from docx import Document
from docx.shared import Pt

ROOT=Path(__file__).resolve().parents[3]
OUT=Path(__file__).resolve().parent
export=runpy.run_path(str(ROOT/'scripts/robust_v2/export_review_docx_20261014.py'))['export']
items=[(ROOT/'paper/manuscript.md',ROOT/'paper/LATEST_MANUSCRIPT_FOR_AVA.docx','AEROFOIL ML | COAUTHOR REVISION | 14 OCTOBER 2026'),(OUT/'ACADEMIC_REFEREE_REPORT.md',OUT/'SCIENTIFIC_REVIEW_REPORT.docx','AEROFOIL ML | ACADEMIC REVIEW | 14 OCTOBER 2026')]
rows=[]
for source,target,header in items:
    export(source,target)
    document=Document(target)
    document.core_properties.title=next(p.text for p in document.paragraphs if p.text.strip())
    document.core_properties.author='Benjamin Ian George Gawith'
    for section in document.sections:
        section.header.paragraphs[0].text=header
        for run in section.header.paragraphs[0].runs:run.font.size=Pt(8)
    for p in document.paragraphs:
        if p.text.startswith(('Table ','Figure ','Supplementary Figure ')):
            p.paragraph_format.keep_with_next=True
        p.paragraph_format.widow_control=True
    for image in document.inline_shapes:
        image._inline.docPr.set('descr', 'Scientific figure; interpretation and limitations are stated in the adjacent caption.')
    document.save(target)
    rows.append({'source':source.relative_to(ROOT).as_posix(),'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'docx':target.relative_to(ROOT).as_posix(),'docx_sha256':hashlib.sha256(target.read_bytes()).hexdigest(),'tables':len(document.tables),'figures':len(document.inline_shapes)})
(OUT/'DOCX_EXPORT.json').write_text(json.dumps(rows,indent=2),encoding='utf-8')
print(json.dumps(rows,indent=2))
